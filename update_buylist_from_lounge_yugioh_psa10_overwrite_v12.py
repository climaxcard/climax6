# -*- coding: utf-8 -*-
"""
トレカラウンジ 遊戯王 PSA10 買取表を取得して、既存 buylist.xlsm のO列だけを直接上書きする。

v13: (25th/片手剣) / (25th/双剣) などの絵柄差分タグをレアリティ扱いで消さず、別商品として照合する。
     C列の括弧内（ホログラフィック/レリーフ/ステンレス等）を判定に使い、Excel側・ラウンジ側の型番が空白でも名前一致でマッチする。

対象Excel（デフォルト）:
  C列 = 商品名
  F列 = 型番/カード番号
  G列 = エキスパンション/レアリティ等（補助）
  O列 = 上書きする価格

特徴:
  - 新しいExcelは作らない。buylist.xlsmを直接保存する。
  - 実行前にバックアップだけ作る。
  - .xlsm のマクロは保持する。keep_vba=True
  - C列の（ホログラフィック）/（レリーフ）等は、名前ではなくレアリティ判定として使う。
  - 型番は QCAC-JP021(QCSE) と QCAC-JP021 の両方を候補として扱う。
  - QCAC-JP084c のような末尾小文字/大文字バリエーションも QCAC-JP084 に寄せる。
  - F列が JP001 など短い場合、G列と結合して QCAC-JP001 のような候補も作る。
  - フル型番は型番一致で更新。ただし同型番でレアリティ違いがある場合はC列/G列のレアリティ指定を必須にする。
  - F列型番が空白の場合は、C列の商品名・括弧タグ（例: ステンレス）でラウンジ商品名と照合する。
  - ラウンジ側の型番が空白の商品も読み捨てず、名前一致用の索引に入れる。
  - ラウンジ価格から5%カットした後、指定単位で切り捨ててO列に入れる。
"""

import argparse
import csv
import re
import shutil
import sys
import time
import unicodedata
from datetime import datetime
from pathlib import Path

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

LOUNGE_URL = "https://kaitori.toreca-lounge.com/products/yugioh/single?grade=PSA10"
DEFAULT_EXCEL = r"C:\Users\user\ClimaxGit\climax6\buylist.xlsm"
DEFAULT_SHEET = None

# 1-indexed columns
NAME_COL = 3       # C
EXPANSION_COL = 7  # G  # 補助。F列が短型番の場合だけ使用
MODEL_COL = 6      # F
PRICE_COL = 15     # O
START_ROW = 6

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/125 Safari/537.36",
    "Accept-Language": "ja,en-US;q=0.9,en;q=0.8",
}

BRACKET_PATTERNS = [
    r"（[^（）]*）", r"\([^()]*\)",
    r"［[^［］]*］", r"\[[^\[\]]*\]",
    r"【[^【】]*】", r"〈[^〈〉]*〉", r"<[^<>]*>",
]

# v7: レリーフ/ホログラフィック等は価格判定に必要なので、名前正規化では消さない。
# 表記ゆれ吸収のために消すのは「新/旧/復刻版」程度に限定。
NOISE_WORDS = [
    "新", "旧", "復刻版", "復刻", "通常",
]


def remove_bracket_text(s):
    if s is None:
        return ""
    s = str(s)
    prev = None
    while prev != s:
        prev = s
        for pat in BRACKET_PATTERNS:
            s = re.sub(pat, "", s)
    return s


def strip_psa10_text(s):
    if s is None:
        return ""
    text = unicodedata.normalize("NFKC", str(s))
    text = re.sub(r"[【\[\(（]?\s*PSA\s*10\s*[】\]\)）]?", "", text, flags=re.IGNORECASE)
    return text


def norm_rarity(s):
    """Excel C列/G列・ラウンジ商品名/レアリティ表記を照合用に正規化。"""
    t = unicodedata.normalize("NFKC", str(s or "")).upper().strip()
    repl = {
        " ": "", "　": "", "・": "", "･": "",
        "－": "-", "‐": "-", "‑": "-", "–": "-", "—": "-", "―": "-", "ー": "-", "ｰ": "-",
        "レア": "", "仕様": "", "カード": "",
    }
    for k, v in repl.items():
        t = t.replace(k, v)
    if any(x in t for x in ["ホログラフィック", "ホロ", "HOLOGRAPHIC", "HR"]):
        return "HOLOGRAPHIC"
    if any(x in t for x in ["アルティメット", "レリーフ", "レリ-フ", "ULTIMATE", "UL"]) or ("レリ" in t and "フ" in t):
        return "ULTIMATE"
    if any(x in t for x in ["クォーターセンチュリーシークレット", "クオーターセンチュリーシークレット", "QCSE", "25TH"]):
        return "QCSE"
    if any(x in t for x in ["プリズマティックシークレット", "プリシク", "PSE"]):
        return "PSE"
    if "20TH" in t:
        return "20TH"
    if "シークレット" in t or t == "SE":
        return "SECRET"
    if "ウルトラ" in t or t == "UR":
        return "ULTRA"
    if "スーパー" in t or t == "SR":
        return "SUPER"
    return t


def rarity_close(excel_rarity, lounge_rarity):
    er = norm_rarity(excel_rarity)
    lr = norm_rarity(lounge_rarity)
    if not er or not lr:
        return False
    return er == lr or er in lr or lr in er


def bracket_contents(s):
    """括弧内テキストを取り出す。例: 【PSA10】(レリーフ)ブラックローズ -> PSA10, レリーフ"""
    if s is None:
        return []
    text = unicodedata.normalize("NFKC", str(s))
    vals = []
    for pat in [r"\(([^()]*)\)", r"（([^（）]*)）", r"\[([^\[\]]*)\]", r"［([^［］]*)］", r"【([^【】]*)】", r"〈([^〈〉]*)〉", r"<([^<>]*)>"]:
        vals.extend([m.strip() for m in re.findall(pat, text) if m and m.strip()])
    return vals


def extract_rarity_from_name(s):
    """C列/ラウンジ商品名から、価格判定に使うレアリティを抽出する。括弧内を最優先。"""
    parts = bracket_contents(s)
    parts.append(strip_psa10_text(s))
    for part in parts:
        r = norm_rarity(part)
        if r in {"HOLOGRAPHIC", "ULTIMATE", "QCSE", "PSE", "20TH", "SECRET", "ULTRA", "SUPER"}:
            return r
    return ""


def strip_rarity_from_name(s):
    """
    名前照合用。PSA10表記と、(ホログラフィック)/(レリーフ)等のレアリティ括弧だけ外す。
    それ以外の括弧（絵柄違い等）は残す。
    """
    if s is None:
        return ""
    text = strip_psa10_text(s)

    def repl(m):
        inner = m.group(1).strip()

        # 片手剣/双剣などの絵柄差分を含む括弧は、
        # 25th/QCSE が含まれていても名前照合から消さない。
        if is_art_variant_tag(inner):
            return m.group(0)

        r = norm_rarity(inner)
        if r in {"HOLOGRAPHIC", "ULTIMATE", "QCSE", "PSE", "20TH", "SECRET", "ULTRA", "SUPER"}:
            return ""
        if re.search(r"PSA\s*10", inner, flags=re.IGNORECASE):
            return ""
        return m.group(0)

    for pat in [r"\(([^()]*)\)", r"（([^（）]*)）", r"\[([^\[\]]*)\]", r"［([^［］]*)］", r"【([^【】]*)】", r"〈([^〈〉]*)〉", r"<([^<>]*)>"]:
        text = re.sub(pat, repl, text)
    return text


def norm_name(s):
    s = strip_rarity_from_name(s)
    s = unicodedata.normalize("NFKC", str(s or "")).upper()
    for w in NOISE_WORDS:
        s = s.replace(w.upper(), "")
    repl = {
        "・": "", "･": "", " ": "", "　": "",
        "'": "", "’": "", "`": "", "´": "",
        "/": "", "／": "", "：": "", ":": "",
        "!": "", "！": "", "?": "", "？": "", "☆": "",
        "－": "", "‐": "", "‑": "", "–": "", "—": "", "―": "", "-": "", "ー": "", "ｰ": "",
        "，": "", ",": "", "。": "", ".": "", "．": "",
        "「": "", "」": "", "『": "", "』": "",
    }
    for k, v in repl.items():
        s = s.replace(k, v)
    s = s.replace("E・HERO", "EHERO").replace("E HERO", "EHERO")
    s = s.replace("D・HERO", "DHERO").replace("D HERO", "DHERO")
    s = s.replace("V・HERO", "VHERO").replace("V HERO", "VHERO")
    s = s.replace("I:P", "IP")
    s = s.replace("LV.", "LV")
    return s.strip()



def norm_name_base_for_tag_match(s):
    """型番空白時の名前照合用。PSA10と全ての括弧を外したベース名。"""
    text = strip_psa10_text(s)
    text = remove_bracket_text(text)
    # norm_name() に渡すと、基本の記号/中黒/長音などを揃えられる
    return norm_name(text)


def is_art_variant_tag(s):
    """片手剣/双剣など、同名・同レアリティでも価格が分かれる絵柄差分タグを判定する。"""
    t = unicodedata.normalize("NFKC", str(s or "")).upper()
    # 区切りや括弧の種類に左右されないように最低限だけ潰す
    for ch in [" ", "　", "・", "･", "-", "－", "‐", "‑", "–", "—", "―", "ー", "ｰ", "/", "／"]:
        t = t.replace(ch, "")
    return any(x in t for x in [
        "片手剣",
        "双剣",
        "イラスト違い",
        "イラスト違",
        "絵違い",
        "絵柄違い",
        "絵柄違",
        "別イラスト",
    ])


def normalize_art_variant_tag(s):
    """絵柄差分タグを照合用キーに正規化する。"""
    t = unicodedata.normalize("NFKC", str(s or "")).upper()
    for ch in [" ", "　", "・", "･", "-", "－", "‐", "‑", "–", "—", "―", "ー", "ｰ", "/", "／", "(", ")", "（", "）", "[", "]", "［", "］", "【", "】"]:
        t = t.replace(ch, "")
    if "片手剣" in t:
        return "ART_KATATEKEN"
    if "双剣" in t:
        return "ART_SOUKEN"
    if any(x in t for x in ["イラスト違い", "イラスト違", "絵違い", "絵柄違い", "絵柄違", "別イラスト"]):
        return "ART_VARIANT"
    return ""


def norm_tag(s):
    """括弧内タグ照合用。ステンレス/レリーフ/ホログラフィック/絵柄差分等を揃える。"""
    # (25th/片手剣) / (25th/双剣) は、25thをQCSE判定しつつも、
    # タグ照合では絵柄差分として分ける。
    art_tag = normalize_art_variant_tag(s)
    if art_tag:
        return art_tag

    r = norm_rarity(s)
    if r in {"HOLOGRAPHIC", "ULTIMATE", "QCSE", "PSE", "20TH", "SECRET", "ULTRA", "SUPER"}:
        return r
    t = unicodedata.normalize("NFKC", str(s or "")).upper()
    repl = {
        " ": "", "　": "", "・": "", "･": "",
        "－": "", "‐": "", "‑": "", "–": "", "—": "", "―": "", "-": "", "ー": "", "ｰ": "",
        "(": "", ")": "", "（": "", "）": "", "[": "", "]": "", "［": "", "］": "", "【": "", "】": "",
        "「": "", "」": "", "『": "", "』": "", "/": "", "／": "",
    }
    for k, v in repl.items():
        t = t.replace(k, v)
    if "ステンレス" in t or "STAINLESS" in t:
        return "STAINLESS"
    if re.search(r"PSA\s*10", t, flags=re.IGNORECASE):
        return ""
    return t.strip()


def name_tags_key(s):
    """商品名の括弧内タグを正規化して tuple 化。PSA10は除外。"""
    vals = []
    for part in bracket_contents(s):
        tag = norm_tag(part)
        if tag and tag not in {"PSA10", "PSA"}:
            vals.append(tag)
    return tuple(sorted(set(vals)))

def norm_code(s):
    if s is None:
        return ""
    s = unicodedata.normalize("NFKC", str(s)).upper().strip()
    s = s.replace("－", "-").replace("‐", "-").replace("‑", "-").replace("–", "-").replace("—", "-").replace("―", "-").replace("ー", "-")
    s = re.sub(r"\s+", "", s)
    return s


def code_variants(code):
    """型番候補を複数作る。QCAC-JP021(QCSE)->QCAC-JP021、QCAC-JP084C->QCAC-JP084 など。"""
    c = norm_code(code)
    if not c or c in {"-", "-/-", "ー", "NONE", "NAN"}:
        return []
    vals = {c}
    vals.add(re.sub(r"\([^)]*\)$", "", c))
    # 末尾レアリティ・絵柄suffixっぽい1文字を外す。例: QCAC-JP084C / QCAC-JP084B
    more = set()
    for v in vals:
        more.add(re.sub(r"([A-Z0-9]+-[A-Z]{1,3}\d{2,3})[A-Z]$", r"\1", v))
        more.add(re.sub(r"([A-Z0-9]+-\d{2,3})[A-Z]$", r"\1", v))
    vals |= more
    return [v for v in vals if v and v not in {"-", "-/-", "ー"}]


def excel_code_candidates(code, expansion=None):
    vals = set(code_variants(code))
    c = norm_code(code)
    exp = norm_code(expansion)
    # FがJP001だけ、GがQCAC-JPなら QCAC-JP001 を候補に追加
    if exp and c and re.fullmatch(r"JP\d{3}[A-Z]?", c) and re.search(r"-[A-Z]{1,3}$", exp):
        prefix = re.sub(r"[A-Z]{1,3}$", "", exp)
        vals.update(code_variants(prefix + c))
    # Fが001だけ、GがQCAC-JPなら QCAC-JP001 を候補に追加
    if exp and c and re.fullmatch(r"\d{3}[A-Z]?", c):
        vals.update(code_variants(exp + c))
    return sorted(vals, key=lambda x: (-len(x), x))


def is_full_code(code):
    c = norm_code(code)
    if not c or c in {"-", "-/-", "ー"}:
        return False
    if "/" in c:
        return True
    return bool(re.search(r"[A-Z0-9]{2,}-[A-Z0-9]{1,}", c))


def parse_lounge_item(text):
    text = " ".join(str(text).split())
    rarity = ""
    m = re.match(r"(.+?)\s+型番\s*[:：]\s*(.*?)\s+レアリティ/種別\s*[:：]\s*(.*?)\s+買取価格\s*[:：]\s*¥?\s*([0-9,]+)", text)
    if m:
        name, code, rarity, price = m.groups()
    else:
        m = re.match(r"(.+?)\s+型番\s*[:：]\s*(.*?)\s+買取価格\s*[:：]\s*¥?\s*([0-9,]+)", text)
        if not m:
            return None
        name, code, price = m.groups()

    # v12: 型番が空白のラウンジ商品も存在するため、ここで捨てない。
    # code_variants が空でも、名前一致用の索引には入れる。
    variants = code_variants(code)

    # ラウンジのシンプル表示では、レアリティ列ではなく商品名の先頭に
    # 「(レリーフ)」「(ホログラフィック)」が入るケースがある。
    # そのため、レアリティ/種別が取れない時は商品名の括弧内から抽出する。
    rarity_from_name = extract_rarity_from_name(name)
    rarity_norm = rarity_from_name or norm_rarity(rarity)
    rarity_label = rarity_from_name or rarity.strip()

    return {
        "name": name.strip(),
        "name_norm": norm_name(name),
        "name_base_norm": norm_name_base_for_tag_match(name),
        "name_tags": name_tags_key(name),
        "rarity": rarity_label,
        "rarity_norm": rarity_norm,
        "code": norm_code(code),
        "code_variants": variants,
        "price": int(price.replace(",", "")),
        "raw": text,
    }


def fetch_lounge(max_pages=10, sleep_sec=0.4):
    session = requests.Session()
    session.headers.update(HEADERS)
    products = []
    seen_raw = set()
    for page in range(1, max_pages + 1):
        url = LOUNGE_URL if page == 1 else f"{LOUNGE_URL}&page={page}"
        print(f"[INFO] ラウンジ取得 page={page}: {url}")
        res = session.get(url, timeout=30)
        res.raise_for_status()
        soup = BeautifulSoup(res.text, "html.parser")
        page_items = []
        # aタグだけを見る。div全体を見ると同じ商品が大量に重複するため。
        for a in soup.find_all("a"):
            text = a.get_text(" ", strip=True)
            if "型番" not in text or "買取価格" not in text:
                continue
            item = parse_lounge_item(text)
            if not item or item["raw"] in seen_raw:
                continue
            seen_raw.add(item["raw"])
            page_items.append(item)
        print(f"[INFO] page={page} 新規取得件数: {len(page_items)}")
        if not page_items:
            break
        products.extend(page_items)
        time.sleep(sleep_sec)
    return products


def build_indexes(products):
    by_name_code = {}
    by_name_code_rarity = {}
    by_code = {}
    by_code_rarity = {}
    by_short_code = {}
    by_name_only = {}
    by_base_tags = {}
    by_base_only = {}
    for p in products:
        # 型番空白行用: 商品名だけでも引ける索引を作る
        by_name_only.setdefault(p.get("name_norm", ""), []).append(p)
        by_base_only.setdefault(p.get("name_base_norm", ""), []).append(p)
        by_base_tags.setdefault((p.get("name_base_norm", ""), p.get("name_tags", tuple())), []).append(p)

        for cv in p.get("code_variants", []):
            by_name_code.setdefault((p["name_norm"], cv), []).append(p)
            by_name_code_rarity.setdefault((p["name_norm"], cv, p.get("rarity_norm", "")), []).append(p)
            if is_full_code(cv):
                by_code.setdefault(cv, []).append(p)
                by_code_rarity.setdefault((cv, p.get("rarity_norm", "")), []).append(p)
            else:
                by_short_code.setdefault(cv, []).append(p)
    return by_name_code, by_name_code_rarity, by_code, by_code_rarity, by_short_code, by_name_only, by_base_tags, by_base_only

def choose_best(cands):
    if not cands:
        return None
    return sorted(cands, key=lambda x: x["price"], reverse=True)[0]


def name_close(excel_name_norm, lounge_name_norm):
    if not excel_name_norm or not lounge_name_norm:
        return False
    if excel_name_norm == lounge_name_norm:
        return True
    if excel_name_norm in lounge_name_norm or lounge_name_norm in excel_name_norm:
        return True
    # 短すぎる名前は緩くしすぎない
    common = set(excel_name_norm) & set(lounge_name_norm)
    denom = max(1, min(len(set(excel_name_norm)), len(set(lounge_name_norm))))
    return len(common) / denom >= 0.82


def find_match(name, code, expansion, by_name_code, by_name_code_rarity, by_code, by_code_rarity, by_short_code, by_name_only, by_base_tags, by_base_only):
    # C列の(ホログラフィック)/(レリーフ)/(ステンレス)等を判定ヒントとして使う
    name_rarity = extract_rarity_from_name(name)
    nname = norm_name(name)
    base_name = norm_name_base_for_tag_match(name)
    tags = name_tags_key(name)
    erarity = name_rarity or norm_rarity(expansion)
    candidates = excel_code_candidates(code, expansion)
    if not nname:
        return None, "empty_name", candidates

    # 0) 型番が空白の場合: 商品名だけで照合する。
    #    オシリスの天空竜(ステンレス) などは、ベース名 + 括弧タグ(STAINLESS)で一致させる。
    if not candidates:
        if base_name and tags:
            cands = by_base_tags.get((base_name, tags), [])
            if erarity:
                rcands = [p for p in cands if rarity_close(erarity, p.get("rarity_norm", ""))]
                if len(rcands) == 1:
                    return rcands[0], "name_base+tags+rarity_no_code", candidates
            if len(cands) == 1:
                return cands[0], "name_base+tags_no_code", candidates
            if len(cands) > 1:
                # 同一タグで複数ある場合は高い方を採用するが、ログで方式が分かるようにする
                return choose_best(cands), "name_base+tags_no_code_multi_best", candidates

        # 括弧の位置がExcelとラウンジで同じ場合の保険
        cands = by_name_only.get(nname, [])
        if len(cands) == 1:
            return cands[0], "name_only_no_code", candidates
        if erarity:
            rcands = [p for p in cands if rarity_close(erarity, p.get("rarity_norm", ""))]
            if len(rcands) == 1:
                return rcands[0], "name_only+rarity_no_code", candidates

        # タグなしでもベース名だけで完全に1件なら更新する。複数なら危険なので未一致にする。
        if base_name:
            cands = by_base_only.get(base_name, [])
            if tags:
                cands = [p for p in cands if set(tags).issubset(set(p.get("name_tags", tuple())))]
            if len(cands) == 1:
                return cands[0], "name_base_only_no_code", candidates
            if erarity:
                rcands = [p for p in cands if rarity_close(erarity, p.get("rarity_norm", ""))]
                if len(rcands) == 1:
                    return rcands[0], "name_base+rarity_no_code", candidates

        return None, "no code and name not unique", candidates

    # 1) 名前+型番+レアリティ。ブラック・ローズ・ドラゴンのホロ/レリーフ等はここで分ける。
    if erarity:
        for c in candidates:
            hit = choose_best(by_name_code_rarity.get((nname, c, erarity), []))
            if hit:
                return hit, "name+code+rarity", candidates

    # 2) フル型番+レアリティ。名前が少し違っても、同型番のレアリティ違いは混ぜない。
    if erarity:
        for c in candidates:
            if is_full_code(c):
                hit = choose_best(by_code_rarity.get((c, erarity), []))
                if hit:
                    return hit, "code+rarity", candidates

    # 3) 名前+型番。ただし同じ名前+型番で複数レアリティがある場合は、C列/G列のレアリティ指定に合うものだけ更新。
    for c in candidates:
        cands = by_name_code.get((nname, c), [])
        if erarity:
            rcands = [p for p in cands if rarity_close(erarity, p.get("rarity_norm", ""))]
            hit = choose_best(rcands)
            if hit:
                return hit, "name+code+rarity_close", candidates
            # レアリティ指定がある行は、レアリティ不一致なら型番だけで混ぜない
            if erarity in {"HOLOGRAPHIC", "ULTIMATE"} or len(cands) > 1:
                continue
        hit = choose_best(cands)
        if hit:
            return hit, "name+code", candidates

    # 4) フル型番のみは最後の保険。複数レアリティがある時は価格混同の原因になるので、単独候補の時だけ更新。
    for c in candidates:
        if is_full_code(c):
            cands = by_code.get(c, [])
            if erarity:
                rcands = [p for p in cands if rarity_close(erarity, p.get("rarity_norm", ""))]
                hit = choose_best(rcands)
                if hit:
                    return hit, "code+rarity_close", candidates
                if erarity in {"HOLOGRAPHIC", "ULTIMATE"}:
                    continue
            if len(cands) == 1:
                hit = cands[0]
                return hit, "code_single", candidates

    # 5) 短い型番は、同型番候補の中で名前とレアリティが近いものだけ。
    for c in candidates:
        if not is_full_code(c):
            cands = [p for p in by_short_code.get(c, []) if name_close(nname, p["name_norm"])]
            if erarity:
                rcands = [p for p in cands if rarity_close(erarity, p.get("rarity_norm", ""))]
                hit = choose_best(rcands)
                if hit:
                    return hit, "short_code+name+rarity", candidates
            hit = choose_best(cands)
            if hit:
                return hit, "short_code+name", candidates

    return None, "no match", candidates


def calc_final_price(price):
    """
    ラウンジ価格から5%カットし、その後O列へ入れる前に切り捨てる。
    例: ラウンジ価格 1,032,000 -> 5%カット 980,400 -> 980,000
        ラウンジ価格   795,000 -> 5%カット 755,250 -> 750,000
        ラウンジ価格    67,500 -> 5%カット  64,125 ->  64,000
        ラウンジ価格     8,900 -> 5%カット   8,455 ->   8,000

    5%カット後の価格に対して、
    100,000円以上: 10,000円単位で切り捨て
    100,000円未満:  1,000円単位で切り捨て
    1,000円未満:     100円単位で切り捨て
    """
    try:
        p = int(str(price).replace(",", "").strip())
    except Exception:
        return price

    # 5%カット（小数は一旦切り捨て）
    p = int(p * 0.95)

    if p >= 100000:
        unit = 10000
    elif p >= 1000:
        unit = 1000
    else:
        unit = 100
    return (p // unit) * unit

def update_excel(excel_path, sheet_name=None, dry_run=False, max_pages=10, start_row=START_ROW, price_col=PRICE_COL):
    excel_path = Path(excel_path)
    if not excel_path.exists():
        raise FileNotFoundError(f"Excelが見つかりません: {excel_path}")

    products = fetch_lounge(max_pages=max_pages)
    if not products:
        raise RuntimeError("ラウンジ商品を取得できませんでした。サイト構造変更の可能性があります。")

    by_name_code, by_name_code_rarity, by_code, by_code_rarity, by_short_code, by_name_only, by_base_tags, by_base_only = build_indexes(products)
    print(f"[INFO] ラウンジ取得合計: {len(products)}件")
    print(f"[INFO] 型番一致キー: {len(by_code)}件 / 型番+レアリティキー: {len(by_code_rarity)}件 / 短型番キー: {len(by_short_code)}件 / 名前+型番キー: {len(by_name_code)}件 / 型番なし名前キー: {len(by_base_only)}件")
    print("[INFO] v12: Excel側/ラウンジ側とも型番空白の商品を、C列の商品名・括弧タグでも照合します。価格は5%カット後に切り捨てます。")

    wb = load_workbook(excel_path, keep_vba=True)
    ws = wb[sheet_name] if sheet_name else wb.active
    print(f"[INFO] 対象シート: {ws.title}")

    log_dir = excel_path.parent / "logs"
    log_dir.mkdir(exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    updated_log = log_dir / f"lounge_yugioh_psa10_updated_{stamp}.csv"
    no_match_log = log_dir / f"lounge_yugioh_psa10_no_match_{stamp}.csv"

    updated = []
    no_match = []
    checked = 0
    sample_rows = []

    for row in range(start_row, ws.max_row + 1):
        name = ws.cell(row=row, column=NAME_COL).value
        expansion = ws.cell(row=row, column=EXPANSION_COL).value
        code = ws.cell(row=row, column=MODEL_COL).value
        if not name and not code:
            continue
        checked += 1
        if len(sample_rows) < 5:
            sample_rows.append((row, name, expansion, code, norm_name(name), extract_rarity_from_name(name) or norm_rarity(expansion), "/".join(excel_code_candidates(code, expansion))))

        hit, method, candidates = find_match(name, code, expansion, by_name_code, by_name_code_rarity, by_code, by_code_rarity, by_short_code, by_name_only, by_base_tags, by_base_only)
        if hit:
            old_price = ws.cell(row=row, column=price_col).value
            raw_price = hit["price"]
            after_5pct = int(int(raw_price) * 0.95)
            new_price = calc_final_price(raw_price)
            ws.cell(row=row, column=price_col).value = new_price
            updated.append([row, name, expansion, code, old_price, raw_price, after_5pct, new_price, hit["name"], hit["code"], hit.get("rarity", ""), method])
        else:
            no_match.append([row, name, expansion, code, norm_name(name), extract_rarity_from_name(name) or norm_rarity(expansion), "/".join(candidates), method])

    print("[INFO] Excel先頭サンプル:")
    for s in sample_rows:
        print(f"  row={s[0]} C={s[1]} / F={s[3]} / G={s[2]} / name_norm={s[4]} / rarity_norm={s[5]} / code候補={s[6]}")

    with updated_log.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["Excel行", "Excel名(C)", "Excel補助(G)", "Excel型番(F)", "旧O列", "ラウンジ元価格", "5%カット後", "最終O列", "ラウンジ名", "ラウンジ型番", "ラウンジレアリティ", "照合方法"])
        w.writerows(updated)

    with no_match_log.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["Excel行", "Excel名(C)", "Excelレアリティ(G)", "Excel型番(F)", "正規化名", "正規化レアリティ", "型番候補", "理由"])
        w.writerows(no_match)

    print(f"[INFO] チェック行数: {checked}")
    print(f"[OK] 更新件数: {len(updated)}")
    print(f"[INFO] 未一致件数: {len(no_match)}")
    print(f"[INFO] 更新ログ: {updated_log}")
    print(f"[INFO] 未一致ログ: {no_match_log}")

    if dry_run:
        print("[DRY-RUN] Excelは保存していません。")
        return

    backup_path = excel_path.with_name(f"{excel_path.stem}_backup_before_lounge_{stamp}{excel_path.suffix}")
    shutil.copy2(excel_path, backup_path)
    print(f"[INFO] バックアップ作成: {backup_path}")
    wb.save(excel_path)
    print(f"[OK] 既存ExcelのO列を直接上書き保存しました: {excel_path}")


def main():
    parser = argparse.ArgumentParser(description="ラウンジ遊戯王PSA10買取表で buylist.xlsm のO列を直接上書き（C列括弧内タグ/F列型番。型番空白は名前一致。ラウンジ型番空白も対応）")
    parser.add_argument("--excel", default=DEFAULT_EXCEL, help="対象buylist.xlsmのパス")
    parser.add_argument("--sheet", default=DEFAULT_SHEET, help="対象シート名。未指定ならアクティブシート")
    parser.add_argument("--dry-run", action="store_true", help="保存せず、更新対象だけ確認")
    parser.add_argument("--max-pages", type=int, default=10, help="取得最大ページ数")
    parser.add_argument("--start-row", type=int, default=START_ROW, help="データ開始行。通常は6")
    parser.add_argument("--price-col", type=int, default=PRICE_COL, help="上書きする列番号。O列=15")
    args = parser.parse_args()
    update_excel(args.excel, sheet_name=args.sheet, dry_run=args.dry_run, max_pages=args.max_pages, start_row=args.start_row, price_col=args.price_col)


if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        print(f"[ERR] {type(e).__name__}: {e}", file=sys.stderr)
        sys.exit(1)
